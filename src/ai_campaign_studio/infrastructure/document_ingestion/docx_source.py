"""DOCX document source parser (S2-G9).

Owns reading a local .docx into per-paragraph ``SourceChunk``s
(``locator_type="docx_para"``, ``locator="p<N>"``). Tables inside the document
body are deliberately NOT extracted in this round (they live outside
``Document.paragraphs`` and require separate ``Document.tables`` traversal) —
documented as OUT_OF_SCOPE for a future follow-up if the need arises.
"""

from __future__ import annotations

import logging

from ai_campaign_studio.domain.common.ids import (
    IngestionRunId,
    SourceChunkId,
    SourceSnapshotId,
)
from ai_campaign_studio.domain.ingestion.entities import SourceChunk
from ai_campaign_studio.infrastructure.document_ingestion.errors import (
    DocumentParseError,
)

logger = logging.getLogger(__name__)


class DocxSource:
    """Extract text chunks from a .docx via python-docx (one per paragraph)."""

    def extract(
        self,
        path: str,
        run_id: IngestionRunId,
        snapshot_id: SourceSnapshotId | None = None,
    ) -> tuple[SourceChunk, ...]:
        if snapshot_id is None:
            raise DocumentParseError(
                "snapshot_id is required — the S2-G6 pipeline must persist the "
                "SourceSnapshot before extracting chunks."
            )
        try:
            from docx import Document  # type: ignore[import-not-found]
        except ImportError as exc:
            raise DocumentParseError(
                "python-docx is not installed; install the 'documents' extra."
            ) from exc

        try:
            document = Document(path)  # type: ignore[call-arg]
        except Exception as exc:
            raise DocumentParseError(
                f"Cannot open DOCX (type={type(exc).__name__})"
            ) from exc

        chunks: list[SourceChunk] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):  # type: ignore[union-attr]
            text = paragraph.text.strip()  # type: ignore[union-attr]
            if not text:
                continue
            locator = f"p{index}"
            chunks.append(
                SourceChunk(
                    id=SourceChunkId(f"{run_id}:{locator}"),
                    snapshot_id=snapshot_id,
                    locator_type="docx_para",
                    locator=locator,
                    text=text,
                )
            )
        return tuple(chunks)


__all__ = ["DocxSource"]
